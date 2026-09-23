"""Дополняет существующий отчёт расчётами, сохраняя ручные правки.

Запуск из корня: python3 scripts/lab_04_vigenere/full_calculations.py
Повторный запуск заменяет только ранее добавленные этим скриптом блоки.
"""
from __future__ import annotations

from collections import Counter
from copy import deepcopy
from fractions import Fraction
from pathlib import Path
import math
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from math_helpers import ALPHABET33, RU_FREQ, columns, decrypt, index_of_coincidence, key_length_scores, recover_key, to_indices
from report_builder import add_heading, add_math, add_page_break, add_para, make_doc, m_frac, m_op, m_sub, m_sup, m_text, omml_display

MARKER = 'vigenere-full-calculations-'
REPORT = ROOT / 'docs/reports/2025/lab_04_vigenere/var_11/Ковалев Д.П. ВКБ43 4 лаба.docx'


def number(value, digits=8):
    return f'{value:.{digits}f}'.replace('.', ',')


def symbol(base, sub):
    return m_sub(m_text(base), m_text(str(sub)))


def equation(doc, *parts):
    """Настоящая формула OMML; размер выбран для длинных численных подстановок."""
    add_math(doc, omml_display(list(parts)))
    p = doc.paragraphs[-1]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    for run in p._p.xpath('.//m:r'):
        props = OxmlElement('w:rPr')
        size = OxmlElement('w:sz')
        size.set(qn('w:val'), '22')
        props.append(size)
        text = run.find(qn('m:t'))
        run.insert(run.index(text), props)


def expanded_integer_sum(doc, label, terms, values, group_size=6):
    """Каждое слагаемое выводится один раз; частичные суммы не округляются."""
    partials = []
    for start in range(0, len(terms), group_size):
        group = terms[start:start+group_size]
        subtotal = sum(values[start:start+group_size])
        partials.append(subtotal)
        equation(doc, symbol(label, len(partials)), m_op(' = '),
                 m_op(' + ').join(group), m_op(f' = {subtotal}'))
    equation(doc, m_text(label), m_op(' = '),
             m_op(' + ').join(m_text(str(v)) for v in partials),
             m_op(f' = {sum(values)}'))
    return sum(values)


def frequency_counts(col):
    counts = Counter(col)
    return [counts[i] for i in range(33)]


def full_ic(doc, col, length, j):
    counts = frequency_counts(col)
    n = len(col)
    add_heading(doc, f'L = {length}, столбец {j}: полный расчёт IC', level=2)
    add_para(doc, f'Берём частоты в порядке алфавита А, Б, В, Г, Д, Е, Ж, З, И, Й, К, Л, М, Н, О, П, Р, С, Т, У, Ф, Х, Ц, Ч, Ш, Щ, Ъ, Ы, Ь, Э, Ю, Я, _. Число символов N = {n}. Для удобства проверки 33 слагаемых числителя разделены на группы по шесть: S₁–S₅ содержат по шесть букв, S₆ — последние три. Ни одно слагаемое не пропущено.')
    terms = [m_text(f'{v}·({v}−1)') for v in counts]
    values = [v*(v-1) for v in counts]
    total = expanded_integer_sum(doc, 'S', terms, values)
    equation(doc, m_text(f'D = N(N−1) = {n}·({n}−1) = {n*(n-1)}'))
    equation(doc, symbol('IC', f'{length};{j}'), m_op(' = '),
             m_frac(m_text(str(total)), m_text(str(n*(n-1)))),
             m_op(f' ≈ {number(total/(n*(n-1)))}'))
    if not math.isclose(total/(n*(n-1)), index_of_coincidence(col), abs_tol=1e-14):
        raise ValueError('Ошибка численной проверки IC')
    return Fraction(total, n*(n-1))


def mean_ic(doc, values, length):
    add_para(doc, 'Среднее вычисляется по неокруглённым дробям выше. Числа в следующей подстановке округлены только для показа; итог взят из точного расчёта.')
    # Не помещаем восемь дробей в одну строку: складываем IC группами по четыре.
    partials = []
    for start in range(0, len(values), 4):
        part = sum(values[start:start+4], Fraction())
        partials.append(part)
        equation(doc, symbol('T', len(partials)), m_op(' ≈ '),
                 m_text(' + '.join(number(float(v)) for v in values[start:start+4])),
                 m_op(f' ≈ {number(float(part))}'))
    result = sum(values, Fraction()) / length
    equation(doc, symbol('ICср', length), m_op(' = '),
             m_frac(m_op(' + ').join(symbol('T', i+1) for i in range(len(partials))), m_text(str(length))),
             m_op(f' ≈ {number(float(result))}'))


def add_main_ic(doc, idx, table):
    add_heading(doc, '3.3. Полные численные подстановки IC для L = 4', level=2)
    add_para(doc, 'Здесь раскрыты суммы для всех четырёх столбцов. Частоты приведены в таблице 6. Полная раскладка 2620 символов по столбцам находится в приложении Г. В приложении Д таким же образом рассчитаны все остальные проверенные длины: 2, 3, 5, 6, 7 и 8. Для IC используется именно n(n−1), а не n²: выбирать одну и ту же позицию дважды нельзя.')
    values = [full_ic(doc, col, 4, j+1) for j, col in enumerate(columns(idx, 4))]
    mean_ic(doc, values, 4)
    add_para(doc, 'Числители четырёх столбцов равны 28984, 29104, 29122 и 29174. Знаменатель одинаков: 655·654 = 428370. Поэтому среднее можно независимо проверить одной дробью:')
    equation(doc, symbol('ICср', 4), m_op(' = '),
             m_frac(m_text('28984 + 29104 + 29122 + 29174'), m_text('4·428370')),
             m_op(' = '), m_frac(m_text('116384'), m_text('1713480')),
             m_op(' ≈ 0,06792259'))


def add_full_columns(doc, cipher, plain, table):
    add_page_break(doc)
    add_heading(doc, 'Приложение Г. Полное разбиение по четырём столбцам')
    add_para(doc, 'Таблица содержит все 655 строк, то есть все 2620 символов шифртекста. Нумерация строк и позиций начинается с 1. В строке r находятся символы с номерами 4r−3, 4r−2, 4r−1 и 4r. Чтобы получить столбец, читаем его сверху вниз. Последнее поле — результат расшифровки данной четвёрки ключом ДВОР; оно приведено для проверки, а не используется при определении IC. Знак «_» учитывается как самостоятельный символ.')
    equation(doc, m_text('N = 2620 / 4 = 655;     j = ((i−1) mod 4) + 1'))
    table(doc, 'Таблица Г.1 — шифртекст по столбцам без сокращений',
          ['Строка', 'Позиции', '1: Д', '2: В', '3: О', '4: Р', 'Открытый блок'],
          [[r+1, f'{4*r+1}–{4*r+4}']+list(cipher[4*r:4*r+4])+[plain[4*r:4*r+4]] for r in range(655)])
    add_para(doc, 'Проверка подсчёта на конкретной букве: выпишем все строки первого столбца, где стоит Г. Каждая из этих строк увеличивает частоту Г на единицу.')
    rows = [r+1 for r in range(655) if cipher[4*r] == 'Г']
    for start in range(0, len(rows), 25):
        add_para(doc, ', '.join(map(str, rows[start:start+25])), indent=False)
    equation(doc, symbol('n', 'Г;1'), m_op(f' = {len(rows)};     '),
             m_text(f'{len(rows)}·({len(rows)}−1) = {len(rows)*(len(rows)-1)}'))


def add_other_lengths(doc, idx, table):
    add_page_break(doc)
    add_heading(doc, 'Приложение Д. Полный расчёт IC для остальных длин')
    add_para(doc, 'Для каждого L текст заново записывается по L символов в строке. Столбец j содержит позиции j, j+L, j+2L и далее до конца текста. Ниже даны начало каждой раскладки, все 33 частоты каждого столбца и полные численные суммы. Раскладка при L = 4 приведена целиком в приложении Г, её четыре IC раскрыты в разделе 3.3.')
    for length in [2, 3, 5, 6, 7, 8]:
        add_heading(doc, f'Д.{length}. Проверка длины L = {length}', level=2)
        quotient, remainder = divmod(len(idx), length)
        equation(doc, m_text(f'2620 = {length}·{quotient} + {remainder}'))
        add_para(doc, f'Первые {remainder} столбцов имеют по {quotient+1} символов, остальные — по {quotient}.' if remainder else f'Во всех {length} столбцах по {quotient} символов.')
        table(doc, f'Таблица Д.{length}.1 — первые восемь строк при L = {length}',
              ['Строка']+[str(j) for j in range(1, length+1)],
              [[r+1]+[ALPHABET33[c] for c in idx[r*length:(r+1)*length]] for r in range(8)])
        cols = columns(idx, length)
        counters = [frequency_counts(col) for col in cols]
        table(doc, f'Таблица Д.{length}.2 — все частоты при L = {length}',
              ['Буква']+[f'Ст. {j}' for j in range(1, length+1)],
              [[a]+[c[i] for c in counters] for i, a in enumerate(ALPHABET33)]+[['Всего']+[len(c) for c in cols]])
        values = [full_ic(doc, col, length, j+1) for j, col in enumerate(cols)]
        mean_ic(doc, values, length)
    add_heading(doc, 'Д.9. Численная проверка выбора длины', level=2)
    scores = key_length_scores(idx, 2, 8)
    peak = max(v for _, v in scores)
    cutoff = peak*0.95
    equation(doc, m_text(f'ICmax ≈ {number(peak)}'))
    equation(doc, m_text(f'Порог = 0,95·ICmax ≈ {number(cutoff)}'))
    for length, ic in scores:
        equation(doc, symbol('ICср', length), m_op(f' ≈ {number(ic)} '+('≥' if ic >= cutoff else '<')+f' {number(cutoff)}'))
    equation(doc, m_text('L = min{4, 8} = 4'))
    add_para(doc, 'Порог сравнивается с полными значениями IC, а не с округлёнными строками таблиц. Длина 8 даёт почти тот же индекс, потому что кратна четырём; выбирается меньшая длина 4.')


def add_mic(doc, idx, table):
    add_page_break(doc)
    add_heading(doc, 'Приложение Е. Подробный подбор ключа через MIC')
    counts = [frequency_counts(c) for c in columns(idx, 4)]
    denominator = 655*655
    add_para(doc, 'Для каждой пары (1, j) вычитаем пробный сдвиг d из букв столбца j. Букве с номером r после этого соответствует исходная частота столбца j в позиции (r+d) mod 33. Числитель MIC — сумма 33 произведений частот. Знаменатель равен N₁·Nⱼ = 655·655, а не 655·654: позиции выбираются из двух разных столбцов.')
    equation(doc, m_text(f'D = 655·655 = {denominator}'))
    numerators = [[sum(counts[0][i]*counts[j][(i+d)%33] for i in range(33)) for d in range(33)] for j in range(1, 4)]
    table(doc, 'Таблица Е.1 — все 33 сдвига для каждой пары: числитель и MIC',
          ['d', 'S₁₂', 'MIC₁₂', 'S₁₃', 'MIC₁₃', 'S₁₄', 'MIC₁₄'],
          [[d]+[item for nums in numerators for item in [nums[d], number(nums[d]/denominator)]] for d in range(33)])
    for j in range(1,4):
        best = max(range(33), key=numerators[j-1].__getitem__)
        for delta in [0, best]:
            add_heading(doc, f'Е.{j}.{delta}. Столбцы 1 и {j+1}, d = {delta}', level=2)
            add_para(doc, 'Произведения идут в порядке букв первого столбца от А до _. Для каждой буквы первый множитель берём из столбца 1, второй — из сдвинутого столбца. Частичные суммы U₁–U₆ служат только для переноса длинного выражения.')
            other = [counts[j][(i+delta)%33] for i in range(33)]
            terms = [m_text(f'{a}·{b}') for a,b in zip(counts[0], other)]
            values = [a*b for a,b in zip(counts[0], other)]
            numerator = expanded_integer_sum(doc, 'U', terms, values)
            equation(doc, symbol('MIC', f'1,{j+1}'), m_text(f'({delta}) = '),
                     m_frac(m_text(str(numerator)), m_text(str(denominator))),
                     m_op(f' ≈ {number(numerator/denominator)}'))
    add_heading(doc, 'Е.4. Все варианты первой буквы ключа', level=2)
    add_para(doc, 'Максимумы MIC дают относительные сдвиги 31, 10, 12. Для каждого t от 0 до 32 строим весь ключ. В таблице приведены все 33 кандидата и первые 16 символов пробной расшифровки. Подчёркивания сохранены.')
    table(doc, 'Таблица Е.2 — перебор общего сдвига', ['t', 'Ключ', 'Первые 16 символов'],
          [[t, ''.join(ALPHABET33[(t+d)%33] for d in [0,31,10,12]), decrypt(idx, ''.join(ALPHABET33[(t+d)%33] for d in [0,31,10,12]))[:16]] for t in range(33)])
    for j, delta in enumerate([0,31,10,12],1):
        equation(doc, symbol('k',j), m_text(f' = (4+{delta}) mod 33 = {(4+delta)%33} → {ALPHABET33[(4+delta)%33]}'))


def add_chi(doc, idx, table):
    add_page_break(doc)
    add_heading(doc, 'Приложение Ж. Полные подстановки χ² для найденных букв')
    add_para(doc, 'Для каждой найденной буквы ключа приведены все 33 наблюдаемых количества O, частоты f, ожидаемые количества E = 655·f и вклады в χ². Затем записана вся сумма из 33 дробей как формулы Word. Для остальных пробных сдвигов итоговые χ² приведены в таблице 12 основного текста; любой из них рассчитывается той же формулой по таблицам частот.')
    key, shifts = recover_key(idx,4)
    for j, (col, (_,shift,expected_chi)) in enumerate(zip(columns(idx,4),shifts),1):
        add_heading(doc, f'Ж.{j}. Столбец {j}, сдвиг {shift}, буква {key[j-1]}', level=2)
        counts = frequency_counts(col)
        terms=[]; contributions=[]; rows=[]
        for i, a in enumerate(ALPHABET33):
            source=(i+shift)%33
            obs=counts[source]
            freq=Fraction(str(RU_FREQ[a]))
            exp=len(col)*freq
            contribution=(obs-exp)**2/exp
            contributions.append(contribution)
            rows.append([a, ALPHABET33[source], obs, number(float(freq),3), number(float(exp),3), number(float(contribution),6)])
            terms.append(m_frac(m_sup(m_text(f'({obs}−{number(float(exp),3)})'),m_text('2')),m_text(number(float(exp),3))))
        table(doc, f'Таблица Ж.{j} — все слагаемые для столбца {j}',
              ['Буква', 'Шифрбуква', 'O', 'f', 'E', '(O−E)²/E'], rows)
        add_para(doc, 'Длинная сумма разбита на группы Q₁–Q₁₁ по три соседние буквы алфавита. Внутри дробей указаны точные значения O и E. Частичные суммы справа округлены до восьми знаков; полный результат рассчитан без промежуточного округления.')
        partials=[]
        for start in range(0,33,3):
            partial=sum(contributions[start:start+3],Fraction())
            partials.append(partial)
            equation(doc, symbol('Q',len(partials)),m_op(' = '),m_op(' + ').join(terms[start:start+3]),m_op(f' ≈ {number(float(partial))}'))
        # Два промежуточных итога позволяют не растягивать формулу за поля страницы.
        for start,end,label in [(0,6,'A'),(6,11,'B')]:
            equation(doc,m_text(label+' = '),m_op(' + ').join(symbol('Q',i+1) for i in range(start,end)),m_op(f' ≈ {number(float(sum(partials[start:end],Fraction())))}'))
        total=sum(contributions,Fraction())
        equation(doc,m_sup(m_text('χ'),m_text('2')),m_text(f'({shift}) = A + B ≈ {number(float(total))}'))
        if not math.isclose(float(total),expected_chi,abs_tol=1e-10):
            raise ValueError('Ошибка численной проверки χ²')


def add_decryption(doc, idx, table):
    add_page_break(doc)
    add_heading(doc, 'Приложение З. Подстановки при расшифровке и обратной проверке')
    key=to_indices('ДВОР')
    add_para(doc, 'Первые 32 позиции расписаны по отдельности: номер шифрбуквы c, номер буквы ключа k, вычитание с прибавлением 33, остаток p и открытая буква. Для каждой позиции сразу показана обратная операция. Остальные позиции вычисляются по тому же правилу; все 655 расшифрованных четвёрок приведены в таблице Г.1.')
    for i,c in enumerate(idx[:32]):
        k=key[i%4]; raw=c-k+33; p=raw%33
        equation(doc,symbol('p',i+1),m_text(f' = ({c}−{k}+33) mod 33 = {raw} mod 33 = {p} → {ALPHABET33[p]}'))
        equation(doc,symbol('c',i+1),m_text(f' = ({p}+{k}) mod 33 = {(p+k)%33} → {ALPHABET33[c]}'))
    plain=[(c-key[i%4]+33)%33 for i,c in enumerate(idx)]
    restored=[(p+key[i%4])%33 for i,p in enumerate(plain)]
    mismatches=sum(a!=b for a,b in zip(idx,restored))
    equation(doc,m_text(f'Проверено позиций: {len(idx)}; несовпадений: {mismatches}'))
    if restored!=idx: raise ValueError('Ошибка обратного шифрования')


def marked_block(fragment, name):
    """Параграфы остаются обычными; закладки ограничивают обновляемые области."""
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(8100+name)); start.set(qn('w:name'),MARKER+str(name))
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(8100+name))
    return [start]+[deepcopy(e) for e in fragment._element.body if e.tag!=qn('w:sectPr')]+[end]


def remove_previous(doc):
    body=doc._element.body
    removing=False
    bookmark_id=None
    for el in list(body):
        if el.tag==qn('w:bookmarkStart') and el.get(qn('w:name'),'').startswith(MARKER):
            removing=True;bookmark_id=el.get(qn('w:id'))
        if removing:
            body.remove(el)
            if el.tag==qn('w:bookmarkEnd') and el.get(qn('w:id'))==bookmark_id:removing=False


def add_full_calculations(doc):
    from detailed_variant_11 import table
    cipher=(ROOT/'artifacts/lab_04_vigenere/cipher_texts/var_11.txt').read_text().strip()
    idx=to_indices(cipher)
    if len(idx)!=2620 or len(cipher)!=len(idx):raise ValueError('Неожиданные исходные данные')
    remove_previous(doc)
    first=make_doc()
    add_main_ic(first,idx,table)
    anchor=next(p._p for p in doc.paragraphs if p.text=='4. Шаг 2. Нахождение букв ключа')
    for elem in marked_block(first,1):anchor.addprevious(elem)
    appendix=make_doc()
    add_full_columns(appendix,cipher,decrypt(idx,'ДВОР'),table)
    add_other_lengths(appendix,idx,table)
    add_mic(appendix,idx,table)
    add_chi(appendix,idx,table)
    add_decryption(appendix,idx,table)
    for elem in marked_block(appendix,2):doc._element.body.insert(len(doc._element.body)-1,elem)
    # Все новые заголовки остаются на одной странице со следующим абзацем.
    for fragment_name in (MARKER+'1',MARKER+'2'):
        active=False
        for element in doc._element.body:
            if element.tag==qn('w:bookmarkStart') and element.get(qn('w:name'))==fragment_name:active=True
            if active and element.tag==qn('w:p') and element.xpath('./w:r/w:rPr/w:b'):
                props=element.get_or_add_pPr()
                if props.find(qn('w:keepNext')) is None:props.append(OxmlElement('w:keepNext'))
            if active and element.tag==qn('w:bookmarkEnd'):break


if __name__=='__main__':
    doc=Document(REPORT)
    add_full_calculations(doc)
    temp=REPORT.with_suffix('.updated.docx')
    doc.save(temp)
    Document(temp)  # Убедиться, что пакет открывается, до замены файла отчёта.
    temp.replace(REPORT)
    print(f'Сохранено: {REPORT}')
    print(f'Таблиц: {len(doc.tables)}; формул Word: {len(doc._element.xpath("//m:oMath"))}')
