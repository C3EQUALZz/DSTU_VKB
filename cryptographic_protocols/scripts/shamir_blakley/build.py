"""Подробный отчёт варианта 11: python3 scripts/shamir_blakley/build.py."""

from __future__ import annotations

from fractions import Fraction
from pathlib import Path
import json
import os
import subprocess
import sys
import tempfile

from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
import report_builder as rb

OUT = ROOT / "docs/reports/2026/shamir_blakley/Ковалев Д.П. ВКБ53 3 лаба.docx"
VARIANTS = json.loads((ROOT / "crates/shamir_blakley/data/variants.json").read_text(encoding="utf-8"))["variants"]
VARIANT_11 = next(variant for variant in VARIANTS if variant["number"] == 11)
X = VARIANT_11["shamir"]["x"]
SHAMIR = VARIANT_11["shamir"]["values"]
A = VARIANT_11["blakley"]["coefficients"]
BLAKLEY = VARIANT_11["blakley"]["values"]
CODES = [197, 213, 200, 196, 205, 192]
WORD = "ЕХИДНА"
W3 = [Fraction(5, 2), Fraction(-4), Fraction(5, 2)]
W5 = [Fraction(130, 33), Fraction(-52, 3), Fraction(130, 7), Fraction(-13, 3), Fraction(12, 77)]
COFACTORS = [[-46, -289, 337], [24, 66, -78], [70, 505, -565]]


def equation(doc, expression: str) -> None:
    """Каждая строка расчёта — редактируемый объект Word OMML."""
    rb.add_math(doc, rb.omml_display([rb.m_text(expression, italic=False)]))


def fraction(doc, numerator: str, denominator: str, *, prefix: str = "", suffix: str = "") -> None:
    rb.add_math(doc, rb.omml_display([
        rb.m_text(prefix, italic=False),
        rb.m_frac(rb.m_text(numerator, italic=False), rb.m_text(denominator, italic=False)),
        rb.m_text(suffix, italic=False),
    ]))


def table(doc, headers: list[str], rows: list[list[str]], *, size: int = 10) -> None:
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    for cell, text in zip(grid.rows[0].cells, headers):
        cell.text = str(text)
    grid.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for values in rows:
        for cell, text in zip(grid.add_row().cells, values):
            cell.text = str(text)
    for row in grid.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.first_line_indent = Cm(0)
                para.paragraph_format.line_spacing = 1
                for run in para.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(size)


def signed_sum(terms: list[int] | list[Fraction]) -> str:
    result = str(terms[0])
    for term in terms[1:]:
        result += f" + {term}" if term >= 0 else f" − {abs(term)}"
    return result


def add_input(doc) -> None:
    rb.add_heading(doc, "2. Исходные доли варианта 11")
    rb.add_para(doc, "Таблица 5.1 задаёт по пять долей для каждого из шести символов X₁–X₆. "
                "Ниже данные переписаны по строкам долей без изменения чисел. Для схемы Шамира "
                "в каждой ячейке указано y при x из первого столбца.")
    table(doc, ["№ доли", "x", "X₁", "X₂", "X₃", "X₄", "X₅", "X₆"], [
        [str(i + 1), str(X[i]), *[str(SHAMIR[j][i]) for j in range(6)]] for i in range(5)
    ])
    rb.add_label(doc, "Таблица 1 — Доли Шамира, пары (x, y)")
    rb.add_para(doc, "В схеме Блэкли коэффициенты a₁, a₂, a₃ стоят в первом столбце, "
                "а каждая ячейка Xⱼ содержит правую часть y уравнения a₁Sⱼ+a₂u+a₃v=y.")
    table(doc, ["№ доли", "(a₁,a₂,a₃)", "X₁", "X₂", "X₃", "X₄", "X₅", "X₆"], [
        [str(i + 1), str(tuple(A[i])), *[str(BLAKLEY[j][i]) for j in range(6)]] for i in range(5)
    ], size=9)
    rb.add_label(doc, "Таблица 2 — Доли Блэкли, четвёрки (a₁, a₂, a₃, y)")


def add_theory(doc) -> None:
    rb.add_heading(doc, "3. Математический метод")
    rb.add_para(doc, "В методичке порог равен k=3, число долей n=5. По трём долям Шамира "
                "восстанавливается квадратичный полином. По пяти долям строится интерполяционный "
                "полином четвёртой степени; для согласованных данных его старшие коэффициенты равны нулю. "
                "Модуль p в условии не задан, поэтому все операции выполняются над рациональными числами точно.")
    equation(doc, "fⱼ(x) = aⱼx² + bⱼx + Sⱼ;  Sⱼ = fⱼ(0)")
    equation(doc, "Lᵢ(0) = ∏ (−xₘ)/(xᵢ − xₘ),  m ≠ i;  Sⱼ = Σ yⱼᵢLᵢ(0)")
    rb.add_para(doc, "Для x = 2, 5, 6 коэффициенты базиса Лагранжа в точке 0:")
    fraction(doc, "(−5)(−6)", "(2−5)(2−6)", prefix="L₁(0) = ", suffix=" = 5/2")
    fraction(doc, "(−2)(−6)", "(5−2)(5−6)", prefix="L₂(0) = ", suffix=" = −4")
    fraction(doc, "(−2)(−5)", "(6−2)(6−5)", prefix="L₃(0) = ", suffix=" = 5/2")
    equation(doc, "Sⱼ(3) = (5/2)y₁ − 4y₂ + (5/2)y₃")
    rb.add_para(doc, "При использовании всех пяти x = 2, 5, 6, 8, 13 произведения Лагранжа "
                "в точке 0 дают веса 130/33, −52/3, 130/7, −13/3, 12/77. Их сумма равна 1.")
    for i, x in enumerate(X):
        others = [other for j, other in enumerate(X) if j != i]
        numerator = "·".join(f"(−{other})" for other in others)
        denominator = "·".join(f"({x}−{other})" for other in others)
        fraction(doc, numerator, denominator, prefix=f"L{i + 1}(0) = ", suffix=f" = {W5[i]}")
    equation(doc, "Sⱼ(5) = (130/33)y₁ − (52/3)y₂ + (130/7)y₃ − (13/3)y₄ + (12/77)y₅")
    rb.add_para(doc, "Для Блэкли секрет — первая координата точки (Sⱼ,u,v). Каждая доля "
                "задаёт плоскость. Первые три плоскости образуют квадратную систему, а пять "
                "плоскостей — переопределённую систему, для которой проверяется отсутствие противоречий.")
    equation(doc, "aᵢ₁Sⱼ + aᵢ₂u + aᵢ₃v = yⱼᵢ")
    equation(doc, "A₃ = [[7,23,1],[27,8,15],[25,10,13]]")
    equation(doc, "D = 7(8·13−15·10) − 23(27·13−15·25) + (27·10−8·25) = 300")
    rb.add_para(doc, "По правилу Крамера числители для трёх координат выражаются "
                "через правые части первых трёх уравнений:")
    equation(doc, "Dₛ = det[[y₁,23,1],[y₂,8,15],[y₃,10,13]]")
    equation(doc, "Dᵤ = det[[7,y₁,1],[27,y₂,15],[25,y₃,13]]")
    equation(doc, "Dᵥ = det[[7,23,y₁],[27,8,y₂],[25,10,y₃]]")
    equation(doc, "Dₛ = −46y₁ − 289y₂ + 337y₃")
    equation(doc, "Dᵤ = 24y₁ + 66y₂ − 78y₃")
    equation(doc, "Dᵥ = 70y₁ + 505y₂ − 565y₃")
    equation(doc, "Sⱼ = Dₛ/300;  u = Dᵤ/300;  v = Dᵥ/300")


def add_shamir_for_symbol(doc, index: int) -> None:
    ys = SHAMIR[index]
    secret = CODES[index]
    rb.add_heading(doc, "Схема Шамира: восстановление из трёх долей", level=2)
    rb.add_para(doc, f"Взяты доли (2,{ys[0]}), (5,{ys[1]}), (6,{ys[2]}). "
                "Коэффициенты Лагранжа уже вычислены выше; подставим каждое значение y.")
    contributions = [W3[i] * ys[i] for i in range(3)]
    table(doc, ["i", "xᵢ", "yᵢ", "Lᵢ(0)", "yᵢLᵢ(0)"], [
        [str(i + 1), str(X[i]), str(ys[i]), str(W3[i]), str(contributions[i])]
        for i in range(3)
    ])
    equation(doc, f"S₍₃₎ = (5/2)·{ys[0]} − 4·{ys[1]} + (5/2)·{ys[2]}")
    equation(doc, f"S₍₃₎ = {signed_sum(contributions)} = {secret}")
    assert sum(contributions) == secret
    delta_1, delta_2 = ys[1] - ys[0], ys[2] - ys[1]
    equation(doc, f"f(5)−f(2) = {ys[1]}−{ys[0]} = {delta_1} = 21a + 3b")
    equation(doc, f"f(6)−f(5) = {ys[2]}−{ys[1]} = {delta_2} = 11a + b")
    equation(doc, f"a = ({3 * delta_2}−{delta_1})/12 = 5;  b = {delta_2}−11·5 = 9")
    equation(doc, f"f₍{index + 1}₎(x) = 5x² + 9x + {secret}")
    rb.add_para(doc, "Проверка исходных трёх долей:")
    for x, y in zip(X[:3], ys[:3]):
        equation(doc, f"f({x}) = 5·{x}² + 9·{x} + {secret} = {y}")

    rb.add_heading(doc, "Схема Шамира: восстановление из пяти долей", level=2)
    rb.add_para(doc, "Теперь используем все пять пар. Последний столбец показывает каждый "
                "слагаемый интерполяционной суммы без округления.")
    contributions = [W5[i] * ys[i] for i in range(5)]
    table(doc, ["i", "xᵢ", "yᵢ", "Lᵢ(0)", "yᵢLᵢ(0)"], [
        [str(i + 1), str(X[i]), str(ys[i]), str(W5[i]), str(contributions[i])]
        for i in range(5)
    ])
    equation(doc, f"S₍₅₎ = (130/33)·{ys[0]} − (52/3)·{ys[1]} + (130/7)·{ys[2]}")
    equation(doc, f"                 − (13/3)·{ys[3]} + (12/77)·{ys[4]}")
    first_part = sum(contributions[:3])
    second_part = sum(contributions[3:])
    equation(doc, f"T₁ = {signed_sum(contributions[:3])} = {first_part}")
    equation(doc, f"T₂ = {signed_sum(contributions[3:])} = {second_part}")
    equation(doc, f"S₍₅₎ = T₁ + T₂ = {signed_sum([first_part, second_part])} = {secret} = S₍₃₎")
    assert sum(contributions) == secret
    rb.add_para(doc, "Полином, интерполированный по пяти долям, имеет коэффициенты "
                f"{secret}, 9, 5, 0, 0 при степенях x⁰–x⁴. "
                "Поэтому он действительно квадратичный; проверяем две дополнительные доли.")
    equation(doc, f"f(8) = 5·8² + 9·8 + {secret} = {ys[3]}")
    equation(doc, f"f(13) = 5·13² + 9·13 + {secret} = {ys[4]}")
    assert all(5 * x * x + 9 * x + secret == y for x, y in zip(X, ys))


def add_blakley_for_symbol(doc, index: int) -> None:
    ys = BLAKLEY[index]
    secret = CODES[index]
    rb.add_heading(doc, "Схема Блэкли: восстановление из трёх долей", level=2)
    rb.add_para(doc, "Составляем три уравнения по первым трём плоскостям:")
    for i in range(3):
        a, b, c = A[i]
        equation(doc, f"{a}S + {b}u + {c}v = {ys[i]}")
    rb.add_para(doc, "Определитель матрицы коэффициентов D=300≠0, значит решение единственно. "
                "Вычисляем три определителя Крамера с подстановкой чисел данного символа:")
    numerators = []
    labels = ["Dₛ", "Dᵤ", "Dᵥ"]
    for label, weights in zip(labels, COFACTORS):
        terms = [weight * y for weight, y in zip(weights, ys[:3])]
        numerator = sum(terms)
        numerators.append(numerator)
        equation(doc, f"{label} = {signed_sum([weights[i] * ys[i] for i in range(3)])} = {numerator}")
    assert numerators == [300 * secret, 1500, 2400]
    equation(doc, f"S = {numerators[0]}/300 = {secret};  u = {numerators[1]}/300 = 5;  v = {numerators[2]}/300 = 8")
    rb.add_para(doc, "Подстановка найденной точки в исходные три плоскости:")
    for i in range(3):
        a, b, c = A[i]
        equation(doc, f"{a}·{secret} + {b}·5 + {c}·8 = {ys[i]}")
        assert a * secret + b * 5 + c * 8 == ys[i]

    rb.add_heading(doc, "Схема Блэкли: восстановление из пяти долей", level=2)
    rb.add_para(doc, "К трём уже решённым уравнениям добавляем ещё две плоскости. "
                "Метод Гаусса–Жордана по всем пяти строкам приводит расширенную матрицу "
                "к трём единичным строкам и двум нулевым строкам без противоречия:")
    equation(doc, f"RREF([A₅|y]) = [[1,0,0|{secret}],[0,1,0|5],[0,0,1|8],[0,0,0|0],[0,0,0|0]]")
    rb.add_para(doc, "Чтобы проверить две дополнительные строки вручную, подставляем "
                "полученные S, u и v отдельно в каждое уравнение:")
    for i in (3, 4):
        a, b, c = A[i]
        equation(doc, f"{a}·{secret} + {b}·5 + {c}·8 = {a * secret} + {b * 5} + {c * 8} = {ys[i]}")
        assert a * secret + b * 5 + c * 8 == ys[i]
    equation(doc, f"S₍₅₎ = {secret} = S₍₃₎;  (S,u,v) = ({secret},5,8)")


def add_manual(doc) -> None:
    rb.add_heading(doc, "4. Подробное ручное восстановление каждого символа")
    rb.add_para(doc, "Для каждой позиции ниже отдельно показаны вычисления Шамира по трём и "
                "пяти долям, затем Блэкли по трём и пяти долям. Общие коэффициенты Лагранжа "
                "и Крамера выведены в разделе 3, а все числовые подстановки приведены здесь.")
    for index, (code, letter) in enumerate(zip(CODES, WORD)):
        rb.add_heading(doc, f"4.{index + 1}. Символ X{index + 1}", level=2)
        add_shamir_for_symbol(doc, index)
        add_blakley_for_symbol(doc, index)
        rb.add_para(doc, f"Итог для X{index + 1}: код {code} по обеим схемам и обеим выборкам; "
                    f"по приложению Б это буква «{letter}».")


def run_cli(share_count: int) -> tuple[str, str]:
    command = ["cargo", "run", "--quiet", "-p", "shamir_blakley", "--",
               "--variant", "11", "--method", "all", "--shares", str(share_count)]
    completed = subprocess.run(command, cwd=ROOT, env={**os.environ, "RUST_LOG": "off"},
                               text=True, capture_output=True, check=True)
    output = completed.stdout.strip()
    assert output.count("Восстановленное слово: ЕХИДНА") == 2
    return " ".join(command), output


def terminal_image(command: str, output: str, path: Path) -> None:
    content = f"$ {command}\n{output}"
    subprocess.run([
        "magick", "-background", "#111827", "-fill", "#f9fafb",
        "-font", "/System/Library/Fonts/Menlo.ttc", "-pointsize", "17", "-size", "1350x",
        "caption:" + content, "-bordercolor", "#111827", "-border", "24", str(path),
    ], check=True, capture_output=True)


def add_results(doc) -> None:
    rb.add_heading(doc, "5. Разделённое и восстановленное сообщение")
    table(doc, ["Позиция", "X₁", "X₂", "X₃", "X₄", "X₅", "X₆"], [
        ["Код", *map(str, CODES)],
        ["Буква", *WORD],
    ])
    rb.add_para(doc, "При всех четырёх сочетаниях метода и числа долей получено одно слово: "
                "ЕХИДНА. Русские коды 192–223 из приложения Б соответствуют заглавным буквам "
                "Windows-1251. Название «ASCII» в методичке для этого блока условно.")


def add_run_and_code(doc) -> None:
    rb.add_heading(doc, "6. Программа, запуск и проверка")
    rb.add_para(doc, "Crate shamir_blakley входит в Cargo workspace. Доменный слой решает "
                "системы методом Гаусса–Жордана с рациональными числами произвольной длины. "
                "Слой данных хранит все 20 вариантов таблицы 5.1; данный отчёт использует "
                "только вариант 11. CLI выбирает вариант, метод и число долей; значение "
                "по умолчанию запускает все четыре сочетания для варианта 11.")
    with tempfile.TemporaryDirectory(prefix="shamir_blakley_report_") as directory:
        for count in (3, 5):
            command, output = run_cli(count)
            image_path = Path(directory) / f"run_{count}.png"
            terminal_image(command, output, image_path)
            rb.add_heading(doc, f"Запуск по {count} долям", level=2)
            rb.add_para(doc, "Ниже показан вывод фактического запуска команды. Для компактности "
                        "в снимке отключён диагностический журнал RUST_LOG.")
            doc.add_picture(str(image_path), width=Inches(6.1))
            rb.add_label(doc, f"Рисунок {1 if count == 3 else 2} — запуск для {count} долей")
            rb.add_listing(doc, "$ " + command + "\n" + output)
    rb.add_para(doc, "Автоматические тесты проверяют результат всех четырёх запусков, "
                "точность рациональной арифметики, восстановление произвольного квадратичного "
                "полинома, отказ при противоречивой пятой доле и совпадающих x. "
                "Общая проверка репозитория выполняется командой just ci.")
    rb.add_heading(doc, "7. Листинг ключевого Rust-кода")
    for title, relative in [
        ("Восстановление Шамира", "crates/shamir_blakley/src/domain/shamir.rs"),
        ("Восстановление Блэкли", "crates/shamir_blakley/src/domain/blakley.rs"),
        ("Точное решение системы", "crates/shamir_blakley/src/domain/linear.rs"),
    ]:
        rb.add_heading(doc, title, level=2)
        source = (ROOT / relative).read_text(encoding="utf-8").split("#[cfg(test)]", 1)[0]
        rb.add_listing(doc, source)


def add_questions_and_conclusion(doc) -> None:
    rb.add_heading(doc, "8. Контрольные вопросы")
    rb.add_qa(doc, 1, "Что такое разбиение данных?", "Разбиение данных — представление "
              "исходного сообщения несколькими фрагментами. Например, слово ЕХИДНА "
              "можно разбить на ЕХИ и ДНА: для составления полного слова нужны оба "
              "фрагмента. Пороговая схема расширяет эту идею возможностью восстановить "
              "секрет не по всем выданным долям.")
    rb.add_qa(doc, 2, "Что такое пороговое разделение?", "В схеме (3,5) "
              "распространяют пять долей одного секрета, а для восстановления требуется "
              "не менее трёх согласованных долей. В варианте 11 первые три доли "
              "восстанавливают код 197 первого символа; все пять подтверждают его.")
    rb.add_qa(doc, 3, "Каков принцип схемы Шамира?", "Секрет является свободным "
              "членом полинома степени k−1. Каждая доля — точка (x,f(x)); "
              "по k точкам полином восстанавливается интерполяцией. Здесь k=3 и f(0)=S.")
    rb.add_qa(doc, 4, "Каков принцип схемы Блэкли?", "Секрет — координата "
              "общей точки гиперплоскостей. Каждая доля задаёт одну плоскость; "
              "пересечение трёх независимых плоскостей определяет точку (S,u,v).")
    rb.add_heading(doc, "9. Выводы")
    rb.add_para(doc, "Вариант 11 выполнен по обеим схемам при использовании трёх и пяти "
                "долей. Во всех случаях восстановлены коды 197, 213, 200, 196, 205, 192 "
                "и сообщение «ЕХИДНА». Каждая из пяти долей каждого символа согласована "
                "с полиномом 5x²+9x+S и точкой (S,5,8). Пять долей позволяют дополнительно "
                "проверить правильность исходных данных. Программа использует точную "
                "арифметику; вычисления в отчёте повторяют её результат.")
    rb.add_para(doc, "В этой методичке схема Шамира задана без конечного поля. Поэтому "
                "показанная численная модель служит учебному восстановлению долей; "
                "криптографическую стойкость стандартной схемы Шамира из неё не следует выводить.")


def main() -> None:
    doc = rb.make_doc()
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
    rb.add_title_page(doc, rb.LabMeta(
        number=3,
        title="Пороговое разделение данных по схемам Шамира и Блэкли",
        variant=11,
        year=2026,
        teacher="Дубровина А.С.",
        student_group="ВКБ53",
    ))
    for paragraph in doc.paragraphs:
        if paragraph.text == "Проверил:":
            paragraph.runs[0].text = "Проверила:"
    rb.add_heading(doc, "1. Тема, цель и задание")
    doc.paragraphs[-1].paragraph_format.page_break_before = True
    rb.add_para(doc, "Тема: математические модели порогового разделения данных по схемам "
                "Шамира и Блэкли. Цель: восстановить сообщение из долей, сравнить "
                "восстановление по трём и пяти долям и реализовать точную проверку на Rust.")
    rb.add_para(doc, "По расписанию это лабораторная работа №3. Исходный текст задания "
                "преподаватель взял из лабораторной работы №5 методических указаний "
                "«Персональная кибербезопасность», страницы 109–125. Использован "
                "только вариант 11 таблицы 5.1 на странице 120.")
    rb.add_para(doc, "Требуется при k=3 и n=5 восстановить закодированное слово "
                "схемами Шамира и Блэкли: отдельно по трём и по пяти долям. "
                "Расчёты всех шести символов и обеих схем приведены ниже.")
    add_input(doc)
    add_theory(doc)
    add_manual(doc)
    add_results(doc)
    add_run_and_code(doc)
    add_questions_and_conclusion(doc)
    for paragraph in doc.paragraphs:
        if paragraph.runs and all(run.bold for run in paragraph.runs):
            paragraph.paragraph_format.keep_with_next = True
    rb.save(doc, OUT)
    print(OUT)


if __name__ == "__main__":
    main()
