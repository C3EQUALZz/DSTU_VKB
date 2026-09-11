"""Собрать DOCX-отчёт из исходников и снимков работающего приложения."""

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Ковалев Д.П. ВКБ53 ЛР1.docx"


def paragraph(doc: DocumentType, text: str, *, center: bool = False) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if center:
        p.paragraph_format.first_line_indent = Cm(0)


def heading(doc: DocumentType, text: str, *, new_page: bool = True) -> None:
    p = doc.add_heading(text, level=1)
    p.paragraph_format.page_break_before = new_page


def code(doc: DocumentType, text: str, caption: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text.rstrip())
    run.font.name = "Courier New"
    run.font.size = Pt(8)
    paragraph(doc, caption, center=True)


def source(doc: DocumentType, path: str, caption: str) -> None:
    code(doc, (ROOT / path).read_text(), caption)


def picture(doc: DocumentType, name: str, caption: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "docs" / "screenshots" / name), width=Cm(16))
    paragraph(doc, caption, center=True)


def main() -> None:
    doc = Document()
    doc.core_properties.author = "Ковалев Данил Петрович"
    doc.core_properties.title = "ЛР1. Создание Web-приложения с использованием Flask"
    doc.core_properties.subject = "Разработка защищённых Web-ресурсов"
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name, style.font.size = "Times New Roman", Pt(14)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(12)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    for text in (
        "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ",
        (
            "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ БЮДЖЕТНОЕ ОБРАЗОВАТЕЛЬНОЕ "
            "УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ"
        ),
        "«ДОНСКОЙ ГОСУДАРСТВЕННЫЙ ТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ» (ДГТУ)",
        "",
        "Факультет «Информатика и вычислительная техника»",
        "Кафедра «Кибербезопасность информационных систем»",
        "",
        "",
        "Лабораторная работа № 1",
        "по дисциплине «Разработка защищённых Web-ресурсов»",
        (
            "на тему «Создание Web-приложения с использованием фреймворка "
            "Flask для расчёта теоретической температуры горения»"
        ),
        "",
    ):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.space_after = Pt(10)
        if text == "Лабораторная работа № 1":
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(16)
    for text in (
        "Выполнил обучающийся гр. ВКБ53",
        "Ковалев Данил Петрович",
        "",
        "Проверил: Ковальчик Р. В.",
    ):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1
    paragraph(doc, "", center=True)
    paragraph(doc, "Ростов-на-Дону, 2026", center=True)

    heading(doc, "1. Тема, цель и условие задания")
    paragraph(
        doc,
        (
            "Тема: создание начального Flask-приложения для технологических"
            " расчётов доменной плавки."
        ),
    )
    paragraph(
        doc,
        (
            "Цель: освоить создание изолированного Python-окружения, "
            "организацию проекта, маршрутизацию Flask и передачу "
            "HTML-шаблона браузеру."
        ),
    )
    paragraph(
        doc,
        (
            "По методическим указаниям требуется создать проект, установить"
            " Flask, зафиксировать зависимости, подготовить index.html и "
            "tempr.html, реализовать GET / и проверить домашнюю страницу в "
            "браузере."
        ),
    )
    paragraph(
        doc,
        (
            "Домашняя страница должна содержать заголовок и меню внутри "
            "header/nav, справочную таблицу внутри main и копирайт с "
            "автором внутри footer. Таблица переносится с рисунка 2 "
            "методических указаний."
        ),
    )
    paragraph(
        doc,
        (
            "По дополнительным требованиям проект создан только через uv, "
            "используется src-структура и чистая архитектура по образцу "
            "answer_service, зависимости внедряются через Dishka. Добавлен "
            "запуск в Docker через Gunicorn."
        ),
    )
    paragraph(
        doc,
        (
            "Граница первой работы: готовая домашняя страница и заготовка "
            "/tempr. В предоставленном задании форма и формула расчёта "
            "температуры отнесены ко второй части, поэтому численный расчёт"
            " здесь не реализуется."
        ),
    )
    doc.add_heading("Краткая теория", level=2)
    paragraph(
        doc,
        (
            "Flask связывает HTTP-путь и метод с Python-обработчиком. "
            "Обработчик получает справочные данные и вызывает "
            "render_template; Jinja формирует HTML. Браузер получает HTML и"
            " отдельно загружает CSS. Чистая архитектура отделяет "
            "предметную модель от HTTP и сборки приложения."
        ),
    )

    heading(doc, "2. Окружение и структура проекта")
    paragraph(
        doc,
        (
            "Проект и окружение .venv созданы менеджером uv. Использованы "
            "Python 3.14, Flask 3.1.3, Dishka 1.10.1 и Gunicorn 25.3.0. "
            "Точные версии и хеши фиксирует uv.lock; requirements.txt "
            "экспортируется из него для соответствия методичке."
        ),
    )
    code(
        doc,
        (
            "uv init --app --package --name blast-furnace --python 3.14 \\\n"
            "  --build-backend hatchling --vcs none .\n"
            "uv sync --frozen\n"
            "uv run blast-furnace run --port 5001\n"
            "uv export --frozen --no-dev --no-emit-project \\\n"
            "  --format requirements-txt --output-file requirements.txt"
        ),
        "Листинг 1 — создание, воспроизведение окружения и запуск",
    )
    code(
        doc,
        (
            "src/blast_furnace/\n"
            "  app.py\n"
            "  domain/reference.py\n"
            "  application/get_blast_reference.py\n"
            "  infrastructure/__init__.py\n"
            "  presentation/http/\n"
            "    routes.py\n"
            "    templates/{base,index,tempr}.html\n"
            "    static/styles.css\n"
            "  setup/{app_factory,config,ioc,cli}.py\n"
            "tests/{unit,integration}/\n"
            "deploy/Dockerfile\n"
            "docker-compose.yaml\n"
            "docker-compose.dev.yaml\n"
            "pyproject.toml, uv.lock, requirements.txt\n"
            "justfile, .importlinter"
        ),
        "Листинг 2 — структура проекта",
    )
    paragraph(
        doc,
        (
            "Внутренние слои domain и application не импортируют Flask, "
            "Dishka или внешние адаптеры. Presentation вызывает прикладной "
            "сценарий. Setup создаёт приложение и контейнер зависимостей. "
            "Infrastructure пока не содержит адаптеров: хранение данных и "
            "внешние системы в ЛР1 отсутствуют."
        ),
    )
    paragraph(
        doc,
        (
            "Границы слоёв контролируются четырьмя контрактами "
            "import-linter. Корневой app.py перенесён внутрь пакета по "
            "требованию к src-структуре. Проект размещён в существующем "
            "Git-репозитории DSTU_VKB."
        ),
    )

    heading(doc, "3. Сборка приложения и внедрение зависимостей")
    source(
        doc,
        "src/blast_furnace/setup/app_factory.py",
        "Листинг 3 — фабрика Flask-приложения",
    )
    source(
        doc,
        "src/blast_furnace/setup/ioc.py",
        "Листинг 4 — регистрация прикладного сценария в Dishka",
    )
    paragraph(
        doc,
        (
            "Фабрика регистрирует Blueprint и интеграцию Dishka. Для "
            "каждого HTTP-запроса создаётся REQUEST-область; обработчик "
            "получает GetBlastReference через FromDishka. Настройки автора "
            "и года передаются шаблонам через конфигурацию Flask."
        ),
    )
    source(doc, "src/blast_furnace/app.py", "Листинг 5 — WSGI-точка входа пакета")

    heading(doc, "4. Маршруты и справочная модель")
    source(
        doc,
        "src/blast_furnace/presentation/http/routes.py",
        "Листинг 6 — обработчики GET / и GET /tempr",
    )
    paragraph(
        doc,
        (
            "Справочная модель BlastParameterEffect — неизменяемый "
            "dataclass. Поля: показатель, изменение параметра, изменение "
            "расхода кокса и изменение производительности. Для чисел "
            "применяется Decimal, а десятичная запятая формируется фильтром"
            " представления decimal_ru."
        ),
    )
    code(
        doc,
        (
            "Температура дутья       | 100 | 0,5 | 1,0\n"
            "Горячая прочность кокса |   1 |   1 | 1,6\n"
            "Расход природного газа |  10 | 0,5 | 1,0"
        ),
        "Листинг 7 — значения справочника с рисунка 2 методички",
    )
    paragraph(
        doc,
        (
            "В исходной таблице не указаны единицы измерения и направления "
            "изменений. Они не добавляются предположительно, а сами "
            "справочные значения не используются как формула расчёта "
            "температуры."
        ),
    )

    heading(doc, "5. Домашняя страница")
    paragraph(
        doc,
        (
            "Общий шаблон base.html содержит семантические блоки header, "
            "nav, main и footer. Шаблон index.html расширяет его и выводит "
            "строки справочника циклом Jinja. Заголовки таблицы имеют "
            "scope, таблица — caption. На узком экране прокручивается "
            "только таблица."
        ),
    )
    picture(
        doc, "index-desktop.png", "Рисунок 1 — главная страница работающего приложения"
    )
    paragraph(
        doc,
        (
            "Страница открыта в Chrome по адресу http://127.0.0.1:5001 из "
            "работающего Docker-контейнера. Внизу отображаются текущий год "
            "и автор: Ковалев Данил Петрович."
        ),
    )

    heading(doc, "6. Страница расчёта температуры")
    paragraph(
        doc,
        (
            "Ссылка «Расчёт ТТГ» ведёт на /tempr. Страница явно сообщает, "
            "что форма и алгоритм будут добавлены во второй части, и "
            "содержит ссылку возврата к справочнику."
        ),
    )
    picture(doc, "tempr-desktop.png", "Рисунок 2 — заготовка страницы расчёта ТТГ")
    source(
        doc,
        "src/blast_furnace/application/get_blast_reference.py",
        "Листинг 8 — прикладной сценарий чтения справочника",
    )
    paragraph(
        doc,
        (
            "Автоматическая браузерная проверка дополнительно выполнена при"
            " ширине 390 пикселей: страница не выходит за пределы окна. "
            "Снимок мобильной версии сохранён в "
            "docs/screenshots/index-mobile.png."
        ),
    )

    heading(doc, "7. Запуск в Docker")
    paragraph(
        doc,
        (
            "Образ собирается в несколько стадий. В builder uv "
            "устанавливает зависимости по lock-файлу и пакет с HTML/CSS. В "
            "runtime копируется готовое окружение без инструментов "
            "тестирования. Приложение запускается от пользователя app (UID "
            "10001)."
        ),
    )
    source(doc, "deploy/Dockerfile", "Листинг 9 — Dockerfile")
    code(
        doc,
        (
            "docker compose up --build -d --wait\n"
            "docker compose logs -f app\n"
            "docker compose down"
        ),
        "Листинг 10 — запуск, журнал и остановка контейнера",
    )
    paragraph(
        doc,
        (
            "Compose публикует порт 127.0.0.1:5001 → 8080. Работают два "
            "процесса Gunicorn; healthcheck запрашивает главную страницу. "
            "Файловая система контейнера доступна только для чтения, /tmp "
            "выделен отдельно. Для разработки применяется overlay "
            "docker-compose.dev.yaml с исходниками и явным --debug."
        ),
    )

    heading(doc, "8. Проверка результата")
    paragraph(
        doc,
        (
            "Проверены форматирование и стиль Ruff, строгая типизация mypy,"
            " статический анализ Bandit, архитектурные контракты "
            "import-linter и функциональные тесты pytest. Проверки "
            "завершились успешно."
        ),
    )
    code(
        doc,
        (
            "uv run ruff format --check .\n"
            "uv run ruff check .\n"
            "uv run mypy\n"
            "uv run bandit -c pyproject.toml -r src\n"
            "uv run lint-imports\n"
            "uv run pytest --cov --cov-report=term-missing"
        ),
        "Листинг 11 — команды проверки",
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Проверка"
    table.rows[0].cells[1].text = "Результат"
    for check, result in (
        ("pytest", "15 тестов пройдено"),
        ("Ruff и mypy", "Ошибок не обнаружено"),
        ("Bandit", "Замечаний не обнаружено"),
        ("import-linter", "4 контракта соблюдены"),
        ("Docker / Gunicorn", "Контейнер healthy; UID 10001"),
        ("Chrome", "Главная, переход /tempr, экран 390 px"),
    ):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = check, result
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1
                for run in p.runs:
                    run.font.size = Pt(12)
    paragraph(
        doc,
        (
            "HTTP-тесты проверяют точное совпадение всех трёх строк таблицы"
            " с заданием, разметку и копирайт, доступность ссылок и CSS, "
            "заготовку /tempr, ответы 404 и 405, экранирование имени "
            "автора, изоляцию настроек и отсутствие debug по умолчанию."
        ),
    )
    paragraph(
        doc,
        (
            "Текстовый протокол проверок сохранён в docs/verification.txt, "
            "снимки страниц — в docs/screenshots. Полные исходные тексты "
            "находятся в src/blast_furnace."
        ),
    )

    heading(doc, "9. Ответы на контрольные вопросы")
    for question, answer in (
        (
            "1. Что такое виртуальное окружение? Зачем requirements.txt?",
            (
                "Окружение изолирует зависимости проекта и предотвращает конфликты "
                "версий. requirements.txt фиксирует пакеты для переноса и анализа "
                "состава ПО. Здесь он экспортируется через uv из uv.lock."
            ),
        ),
        (
            "2. Чем templates отличается от пакета services?",
            (
                "templates содержит HTML-файлы для Flask/Jinja, поэтому импорт "
                "Python и __init__.py не нужны. services содержит импортируемые "
                "Python-модули; __init__.py обозначает обычный пакет. В проекте "
                "используются пакеты domain и application."
            ),
        ),
        (
            "3. Зачем выносить бизнес-логику в отдельный пакет?",
            (
                "Три причины: тестирование без HTTP-сервера; повторное "
                "использование алгоритма, например из CLI; независимость "
                "предметных правил от изменений шаблонов и фреймворка."
            ),
        ),
        (
            "4. Как связываются URL и обработчик? Что делает render_template?",
            (
                "Декоратор route или get связывает URL и HTTP-метод с функцией. "
                "render_template находит шаблон в настроенном каталоге templates, "
                "подставляет данные средствами Jinja и возвращает HTML."
            ),
        ),
        (
            "5. Что даёт debug=True? Почему встроенный сервер не промышленный?",
            (
                "Debug включает перезагрузку и подробную диагностику с "
                "интерактивным отладчиком. Встроенный сервер предназначен для "
                "разработки, а не для обеспечения промышленной нагрузки и "
                "устойчивости. Отладчик нельзя публиковать для посторонних. "
                "В обычном Docker-режиме используется Gunicorn без debug."
            ),
        ),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        p.add_run(question).bold = True
        paragraph(doc, answer)

    heading(doc, "10. Вывод и использованные источники")
    paragraph(
        doc,
        (
            "Выполнена первая лабораторная работа: создан Python-проект с "
            "окружением uv, реализованы Flask-приложение, домашняя страница"
            " со справочной таблицей и страница-заготовка расчёта "
            "температуры. Обеспечены семантическая разметка, src-структура,"
            " разделение слоёв и внедрение зависимостей через Dishka."
        ),
    )
    paragraph(
        doc,
        (
            "Подготовлены воспроизводимая контейнерная сборка, запуск через"
            " Gunicorn и режим разработки. Работоспособность подтверждена "
            "автоматическими тестами, статическими проверками, healthcheck "
            "контейнера и открытием страниц в браузере. Проект готов к "
            "расширению в последующих лабораторных работах."
        ),
    )
    for text in (
        (
            "1. ДГТУ. Лабораторная работа № 1 «Создание Web-приложения с "
            "использованием фреймворка Flask для расчёта теоретической "
            "температуры горения». Ростов-на-Дону, 2026. Предоставленный "
            "PDF, с. 3–9."
        ),
        (
            "2. Flask. Application Setup. "
            "https://flask.palletsprojects.com/en/stable/tutorial/factory/"
        ),
        (
            "3. Dishka. Flask integration. "
            "https://dishka.readthedocs.io/en/stable/integrations/flask.html"
        ),
        "4. uv. Using uv in Docker. https://docs.astral.sh/uv/guides/integration/docker/",
        (
            "5. Flask. Gunicorn. "
            "https://flask.palletsprojects.com/en/stable/deploying/gunicorn/"
        ),
        "Дата обращения к электронным источникам: 11.09.2026.",
    ):
        paragraph(doc, text)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
