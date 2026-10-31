"""Воспроизводимый отчёт практики 2: проверки → реальные запуски → DOCX.

Запуск из любой директории: python3 scripts/practice_02/build.py.
Использует существующий шаблон отчётов и freeze; не меняет старые отчёты.
"""

from __future__ import annotations

import json
import os
import re
import shlex
import struct
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from report_builder import (
    LabMeta, add_bullets, add_heading, add_image, add_listing, add_math,
    add_page_break, add_para, add_qa, add_table_simple, add_title_page, make_doc, save,
)
from snippet_tools import render_code_chunked, render_term_chunked

OUT = ROOT / "docs/reports/practice_02"
EVIDENCE = OUT / "evidence"
SHOTS = ROOT / "docs/screenshots/practice_02"
SNIPS = ROOT / "docs/snippets/practice_02"
BINARY = "./target/debug/practice_02"
ENV = {**os.environ, "RUST_LOG": "off", "NO_COLOR": "1", "RUST_BACKTRACE": "0", "CARGO_TERM_COLOR": "never"}


def command(args: list[str], expected: int = 0) -> subprocess.CompletedProcess:
    result = subprocess.run(args, cwd=ROOT, env=ENV, capture_output=True, text=True, timeout=600)
    if result.returncode != expected:
        raise RuntimeError(f"{shlex.join(args)}: expected {expected}, got {result.returncode}\n{result.stdout}\n{result.stderr}")
    return result


def capture(name: str, cases: list[tuple[list[str], int]], *, max_lines: int = 24,
            summarize_errors: bool = False, compact_curve: bool = False) -> list[Path]:
    print(f"Рендер вывода: {name}", flush=True)
    blocks = []
    displayed = []
    for args, expected in cases:
        result = command([BINARY, *args], expected)
        text = f"$ practice_02 {shlex.join(args)}\n{result.stdout}{result.stderr}"
        if expected:
            text += f"Код завершения: {result.returncode}\n"
        blocks.append(text.rstrip())
        if summarize_errors and expected:
            excerpt = result.stderr.split("\nLocation:", 1)[0].strip()
            displayed.append(f"$ practice_02 {shlex.join(args)}\n{result.stdout}{excerpt}\nКод завершения: {result.returncode}")
        else:
            excerpt = text.rstrip()
            if compact_curve:
                excerpt = "\n".join(line for line in excerpt.splitlines() if not line.startswith("E: "))
            displayed.append(excerpt)
    text = "\n\n".join(blocks) + "\n"
    text = "\n".join(line.rstrip() for line in text.splitlines()) + "\n"
    (EVIDENCE / f"{name}.log").write_text(text, encoding="utf-8")
    return render_term_chunked("\n\n".join(displayed) + "\n", SHOTS, name, max_lines=max_lines, wrap=88)


def reference_verification() -> dict:
    """Независимые Python big integers + modular inverse (extended GCD).

    Rust использует u128/Fermat, здесь pow(denominator,-1,p), что дополнительно
    проверяет переполнения и иной способ вычисления обратного элемента.
    """
    cases = []
    for p, a, b, base in [(17, 2, 2, (5, 1)), (97, 2, 3, (3, 6)), (18446744073709551557, 1, 1, (0, 1))]:
        def add(left, right):
            if left is None:
                return right
            if right is None:
                return left
            x, y = left
            u, v = right
            if x == u and (y + v) % p == 0:
                return None
            slope = ((3 * x * x + a) * pow(2 * y, -1, p) if left == right
                     else (v - y) * pow(u - x, -1, p)) % p
            rx = (slope * slope - x - u) % p
            return rx, (slope * (x - rx) - y) % p

        def multiply(k):
            result = None
            current = base if k >= 0 else (base[0], -base[1] % p)
            k = abs(k)
            while k:
                if k & 1:
                    result = add(result, current)
                current = add(current, current)
                k >>= 1
            return result

        for k in [0, 1, 2, 3, 7, 19, -1, -3, 123456789, -(2**127), 2**127 - 1]:
            expected = multiply(k)
            result = command([BINARY, "--p", str(p), "--a", str(a), "--b", str(b), "mul", "--", str(k), f"{base[0]},{base[1]}"])
            expected_text = "O" if expected is None else f"({expected[0]}, {expected[1]})"
            actual = result.stdout.strip().splitlines()[-1].split(" = ")[-1]
            if actual != expected_text:
                raise AssertionError(f"reference mismatch p={p}, k={k}: {actual} != {expected_text}")
            cases.append({"p": p, "a": a, "b": b, "point": base, "k": str(k), "expected": expected_text, "actual": actual})
    report = {"method": "Python arbitrary precision integers and pow(denominator, -1, p)", "passed": len(cases), "cases": cases}
    (EVIDENCE / "reference.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


def snippet(name: str, file: str, start: str, end: str) -> list[Path]:
    source = (ROOT / file).read_text(encoding="utf-8")
    first = source.index(start)
    last = source.index(end, first)
    excerpt = source[first:last].rstrip()
    (EVIDENCE / f"{name}.rs.txt").write_text(f"// {file}\n{excerpt}\n", encoding="utf-8")
    return render_code_chunked(excerpt, SNIPS, name, max_lines=25, font_size=14)


def heading(doc, title: str, level: int = 1):
    add_heading(doc, title, level=level)
    doc.paragraphs[-1].paragraph_format.keep_with_next = True


def formula(doc, text: str):
    add_math(doc, f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>{escape(text)}</m:t></m:r></m:oMath></m:oMathPara>')


def images(doc, paths: list[Path], caption: str, *, max_width: float = 15.5):
    for index, path in enumerate(paths, 1):
        width, height = struct.unpack(">II", path.read_bytes()[16:24])
        # Ограничиваем высоту, чтобы рисунок и подпись помещались на A4.
        width_cm = min(max_width, 16.5 * width / height)
        suffix = f" (часть {index} из {len(paths)})" if len(paths) > 1 else ""
        add_image(doc, path, caption=caption + suffix, width_cm=width_cm)
        doc.paragraphs[-2].paragraph_format.keep_with_next = True
        doc.paragraphs[-1].paragraph_format.keep_together = True


def table(doc, rows, caption):
    add_table_simple(doc, rows, caption=caption)
    for row in doc.tables[-1].rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def build_document(metadata: dict, shots: dict, snippets: dict) -> Path:
    doc = make_doc()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    doc.core_properties.title = "Практическая работа № 2. Калькулятор на эллиптических кривых"
    doc.core_properties.author = "Ковалев Данил Петрович"
    doc.core_properties.subject = "Программирование алгоритмов защиты информации"
    doc.core_properties.language = "ru-RU"
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.paragraph_format.first_line_indent = Cm(0)
    field = OxmlElement("w:fldSimple")
    field.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr", "PAGE")
    footer._p.append(field)
    add_title_page(doc, LabMeta(number=2, title="Калькулятор на эллиптических кривых", kind="Практическая работа"))
    add_page_break(doc)

    heading(doc, "Тема и цель работы")
    add_para(doc, "Тема: вычисления в группе точек эллиптической кривой над простым полем. Цель — освоить модульную арифметику и операции с точками, исследовать группу и продемонстрировать обмен ключами.")
    heading(doc, "Постановка задачи", 2)
    add_para(doc, "Преподавателем задано название без дополнительного условия и варианта. Принята постановка: создать консольный калькулятор для невырожденных кривых Вейерштрасса над Fp, p > 3.")
    add_bullets(doc, [
        "Проверять параметры кривой и принадлежность точки.",
        "Вычислять −P, P+Q, P−Q, 2P, kP и kP+mQ.",
        "Исследовать точки, порядки, подгруппы и таблицу сложения.",
        "Показывать ECDH и перебор дискретного логарифма.",
        "Проверить особые случаи и оформить результаты.",
    ])
    heading(doc, "Используемые средства", 2)
    table(doc, [
        ["Средство", "Назначение"],
        [metadata["rustc"], "Rust edition 2024; минимальная заявленная версия 1.85"],
        ["clap 4; tracing; color-eyre; thiserror", "CLI, журнал операций, диагностические ошибки"],
        ["cargo test; proptest; Clippy; rustfmt", "Автоматизированная проверка и контроль качества"],
        ["Python; python-docx; freeze", "Независимая сверка и генерация отчёта"],
    ], "Таблица 1 — программные средства")

    add_page_break(doc)
    heading(doc, "Теоретические сведения")
    heading(doc, "Поле и эллиптическая кривая", 2)
    add_para(doc, "Поле Fp состоит из классов вычетов 0, …, p−1 по простому модулю p. Все действия с координатами выполняются по модулю p. Деление на ненулевое d означает умножение на обратный элемент d⁻¹. В данной реализации он вычисляется по малой теореме Ферма: d⁻¹ = d^(p−2) mod p.")
    formula(doc, "E(Fp): y² ≡ x³ + ax + b (mod p), p > 3")
    formula(doc, "Δ = −16(4a³ + 27b²) ≢ 0 (mod p)")
    add_para(doc, "Условие Δ ≠ 0 исключает особые точки кривой. К множеству конечных решений (x,y) добавляется точка на бесконечности O. Именно это множество с законом сложения образует абелеву группу [1, § 2.2.1; 2, § 3].")
    heading(doc, "Сложение и удвоение точек", 2)
    add_para(doc, "Для P=(x₁,y₁), Q=(x₂,y₂), P ≠ Q и x₁ ≠ x₂ применяется секущая. При P=Q и y₁ ≠ 0 применяется касательная. Коэффициент λ вычисляется в поле:")
    formula(doc, "λ = (y₂ − y₁)(x₂ − x₁)⁻¹ mod p  (P ≠ Q)")
    formula(doc, "λ = (3x₁² + a)(2y₁)⁻¹ mod p  (P = Q)")
    formula(doc, "x₃ = λ² − x₁ − x₂ mod p")
    formula(doc, "y₃ = λ(x₁ − x₃) − y₁ mod p")
    add_para(doc, "Особые случаи рассматриваются до деления: P+O=P; O+O=O; −P=(x₁,−y₁ mod p); P+(−P)=O. Если y₁=0, то 2P=O. Вычитание определяется как P−Q=P+(−Q) [1, § 2.2.1].")

    add_page_break(doc)
    heading(doc, "Скалярное умножение и структура группы")
    add_para(doc, "Скалярное умножение kP означает повторное сложение точки. Для k=0 результат равен O, для отрицательного k используется противоположная точка. Двоичный алгоритм double-and-add последовательно обрабатывает биты |k|: для установленного бита прибавляет текущую точку к накопителю и удваивает текущую точку перед переходом к следующему разряду [1, § 2.2.1].")
    formula(doc, "k = Σ bᵢ2ⁱ;  kP = Σ bᵢ(2ⁱP)")
    add_para(doc, "Число групповых операций составляет O(log |k|). В аффинных координатах каждая обычная операция сложения содержит модульное обращение; двоичное возведение в степень для обращения требует O(log p) модульных умножений. Поэтому число модульных умножений при скалярном умножении имеет порядок O(log |k| · log p).")
    add_para(doc, "Число точек #E включает O. Порядок точки ord(P) — минимальное положительное n, при котором nP=O. Для O он равен 1. Подгруппа ⟨P⟩ содержит O,P,…,(n−1)P; по теореме Лагранжа ord(P) делит #E. Порядок точки и число точек всей кривой в общем случае различаются.")
    add_para(doc, "Для перечисления точек сначала строится таблица квадратов y² mod p, затем для каждого x вычисляется x³+ax+b и извлекаются соответствующие y. Использование BTreeMap даёт O(p log p) операций работы с таблицей и O(p) памяти; полный двойной перебор x,y потребовал бы O(p²) проверок.")
    add_para(doc, "Поиск порядка выполняется последовательным сложением до первого O. Дискретный логарифм Q=kP ищется начиная с k=0. Если подгруппа замкнулась, а Q не встретилась, решения нет; если исчерпан лимит, выводится только сообщение об исчерпании лимита. Строится полная таблица сложения для малой группы.")

    add_page_break(doc)
    heading(doc, "Программная реализация")
    add_para(doc, "Отдельный crate practice_02 включён в Cargo workspace. Слой domain содержит чистую математику; application организует вычисления и журналирование; presentation разбирает CLI. Внешних портов и инфраструктурных адаптеров у вычислений нет.")
    add_listing(doc, "crates/practice_02/\n  src/domain/field.rs  — поле и проверка простоты\n  src/domain/curve.rs  — точки и групповые операции\n  src/application/    — сценарии калькулятора\n  src/presentation/   — clap CLI\n  tests/domain.rs     — математические проверки\n  tests/cli.rs        — сквозные проверки команд", "Листинг 1 — структура программы")
    add_para(doc, "Точка представлена перечислением Infinity либо Affine{x,y}. Проверка параметров выполняется при создании Curve, проверка принадлежности — перед каждой публичной операцией над точкой. Числа со знаком на входе нормализуются в 0,…,p−1. Скаляр i128::MIN обрабатывается через unsigned_abs без переполнения знакового типа.")
    add_para(doc, "Модуль хранится в u64. Произведение двух u64 помещается в u128: (2⁶⁴−1)² < 2¹²⁸. Поэтому сложение и умножение перед редукцией выполняются в u128, а вычитание — через i128. Готовые криптографические библиотеки и библиотеки длинной арифметики не используются.")
    images(doc, snippets["field"], "Листинг 2 — безопасная модульная арифметика", max_width=13.5)

    add_page_break(doc)
    heading(doc, "Алгоритм сложения точек")
    add_para(doc, "Внутренний метод add_valid вызывается только после проверки точек или для результатов уже корректных операций. Сначала обрабатываются O и противоположные точки, затем выбирается формула λ. Обратный элемент нуля не вычисляется.")
    images(doc, snippets["addition"], "Листинг 3 — реализация закона сложения")

    add_page_break(doc)
    heading(doc, "Алгоритм скалярного умножения")
    images(doc, snippets["multiply"], "Листинг 4 — двоичный double-and-add")
    add_para(doc, "Проверка простоты p использует детерминированный Miller–Rabin с основаниями 2, 325, 9375, 28178, 450775, 9780504, 1795265022. Этот набор покрывает диапазон n < 2⁶⁴ [3]. Он позволяет проверять большие учебные модули без пробного деления до √p.")

    add_page_break(doc)
    heading(doc, "Ручная проверка контрольного примера")
    add_para(doc, "Выберем E: y²=x³+2x+2 mod 17 и P=(5,1). Проверка принадлежности: 1²=1, а 5³+2·5+2=137≡1 mod 17. Условие невырожденности: 4·2³+27·2²=140≡4, Δ≡−16·4≡4≠0.")
    add_para(doc, "Для удвоения: числитель 3·5²+2=77≡9; знаменатель 2·1=2, обратный элемент 2⁻¹≡9. Тогда λ≡9·9≡13; x₃≡13²−5−5≡6; y₃≡13(5−6)−1≡3. Следовательно, 2P=(6,3).")
    add_para(doc, "Для P+(6,3): λ≡(3−1)/(6−5)≡2, x₃≡4−5−6≡10, y₃≡2(5−10)−1≡6. Поэтому 3P=(10,6). Обратная точка −P=(5,16), а P+(−P)=O.")
    table(doc, [
        ["Вычисление", "Ожидаемый результат"],
        ["2P; 3P; 7P", "(6,3); (10,6); (0,6)"],
        ["−3P; 0P; 19P", "(10,11); O; O"],
        ["2P+3Q, Q=2P", "8P=(13,7)"],
        ["#E; ord(P)", "19; 19"],
        ["ECDH: a=5, b=7", "A=(9,16); B=(0,6); S=35P=16P=(10,11)"],
    ], "Таблица 2 — контрольные значения")

    add_page_break(doc)
    heading(doc, "Сборка и команды запуска")
    add_listing(doc, "cargo build -p practice_02\njust practice-02 --help\njust practice-02 demo\n\njust practice-02 info\njust practice-02 check 5,1\njust practice-02 neg 5,1\njust practice-02 add 5,1 6,3\njust practice-02 sub 5,1 6,3\njust practice-02 double 5,1\njust practice-02 mul 7 5,1\njust practice-02 mul -- -3 5,1\njust practice-02 linear 2 5,1 3 6,3\njust practice-02 inverse 2", "Листинг 5 — арифметические операции")
    add_listing(doc, "just practice-02 points\njust practice-02 count\njust practice-02 order 5,1\njust practice-02 subgroup 5,1\njust practice-02 dlog 5,1 0,6\njust practice-02 ecdh 5,1 5 7\njust practice-02 --p 5 --a 1 --b 0 table", "Листинг 6 — исследование группы и ECDH")
    add_para(doc, "Для произвольной кривой задаются --p, --a, --b. Без этих опций используются p=17,a=2,b=2. Команда demo намеренно показывает фиксированный набор примеров на этой кривой. Координаты записываются x,y, нейтральный элемент — O. RUST_LOG=off отключает журнал, RUST_LOG=debug включает диагностику. Вычисленный результат печатается в stdout, ошибки — в stderr.")

    for name, title, caption, explanation in [
        ("01_arithmetic", "Результаты арифметических операций", "Рисунок 1 — реальные запуски основных операций", "Полученные суммы, разности и кратные совпали с ручными контрольными значениями. Включены отрицательный и нулевой скаляры, а также сумма противоположных точек."),
        ("02_points", "Точки демонстрационной кривой", "Рисунок 2 — все точки E17(2,2)", "В списке 19 элементов, включая O. На этой кривой точка (5,1) имеет порядок 19 и порождает всю группу."),
        ("03_subgroup", "Порядок и циклическая подгруппа", "Рисунок 3 — порядок точки и её кратные", "Каждая из 19 точек встречается в списке кратных ровно один раз, начиная с 0P=O; следующий элемент 19P снова равен O."),
        ("04_table", "Таблица сложения малой группы", "Рисунок 4 — таблица E5(1,0)", "Группа состоит из O,(0,0),(2,0),(3,0). Каждая ненулевая точка имеет порядок 2, а число точек всей кривой равно 4. Этот пример показывает различие двух понятий порядка."),
        ("05_ecdh", "Дискретный логарифм и обмен ECDH", "Рисунок 5 — поиск k и совпадение общих точек", "Для Q=(0,6) получено k=7. При a=5,b=7 обе стороны обмена вычислили одну общую точку (10,11). Перебор дискретного логарифма здесь возможен благодаря малому размеру группы."),
        ("06_errors", "Проверка ошибочных входных данных", "Рисунок 6 — фрагменты фактических сообщений об ошибках", "Составной модуль, вырожденная кривая и арифметика с посторонней точкой завершаются ошибкой. Команда check для посторонней точки корректно возвращает false. Превышение лимита полного перебора обозначается явно. На рисунке показаны сообщения и коды завершения; полная диагностика сохранена в журнале 06_errors.log."),
        ("07_large", "Проверка большого модуля", "Рисунок 7 — арифметика у верхней границы u64", "Обычные операции с точками поддерживают весь диапазон u64 для p. Ограничение p≤10000 относится к полному перечислению, а не к арифметике."),
    ]:
        add_page_break(doc)
        heading(doc, title)
        if name == "05_ecdh":
            formula(doc, "A = aG;  B = bG;  Sₐ = aB = abG = bA = Sᵦ")
            add_para(doc, "Участники используют одну точку G и заданные скаляры a,b. Совпадение общих точек следует из свойств скалярного умножения. В программе выводятся публичные и общие точки; случайная генерация секретов и получение сеансового ключа через KDF не выполняются [2, § 4].")
        add_para(doc, explanation)
        images(doc, shots[name], caption)

    add_page_break(doc)
    heading(doc, "Контроль качества и тестирование")
    add_para(doc, f"Проверено {metadata['checked_at']}. just ci: код 0; форматирование, Clippy и тесты workspace прошли успешно. Практика 2: {metadata['practice_tests']} теста, 0 ошибок.")
    table(doc, [
        ["Группа проверок", "Содержание"],
        ["Ручные векторы", "18 ненулевых кратных G и замыкание 19G=O; пример ECDH"],
        ["Групповые свойства", "Все 19³=6859 троек для ассоциативности; proptest: коммутативность, обратимость, распределительность и композиция скаляров"],
        ["Независимый перебор", "Прямая проверка уравнения двойным циклом на нескольких кривых; граница Хассе и делимость порядков"],
        ["Граничные значения", "u64-модуль около 2⁶⁴; i128::MIN/MAX; ноль, O, y=0; псевдопростые числа"],
        ["CLI", "Все подкоманды; своя кривая; неверные аргументы и ненулевые коды ошибок"],
        ["Независимая сверка", f"{metadata['reference_cases']} вычисления на трёх кривых с Python big integers и другим способом модульного обращения"],
    ], "Таблица 3 — выполненные проверки")
    images(doc, shots["08_quality"], "Рисунок 8 — итог обязательных проверок", max_width=13)
    add_para(doc, "Полные журналы и независимые сверки: docs/reports/practice_02/evidence. Все итоги получены из фактических запусков.")

    add_page_break(doc)
    heading(doc, "Ограничения и обсуждение решения")
    add_bullets(doc, [
        "Поддерживается короткая форма Вейерштрасса над простым полем характеристики >3; расширения F₂ᵐ и стандартные 256-битные кривые не входят в реализацию.",
        "Полное перечисление точек разрешено при p≤10000, таблица — до 32 точек, поиск порядка/подгруппы/логарифма — до 100000 сложений. Ограничения предотвращают длительные переборы и чрезмерный вывод.",
        "Скалярное умножение двоичное, координаты аффинные. Проективные координаты, алгоритм Шуфа и быстрые методы дискретного логарифма не реализованы.",
        "ECDH — учебное вычисление общих точек с заданными скалярами. Постоянное время, случайная генерация ключей, KDF и прикладное шифрование отсутствуют. Код не предназначен для защиты реальных данных.",
    ])
    add_para(doc, "Умножение двух точек и обычное деление точки на точку не являются стандартными групповыми операциями. Выражение kP означает умножение на целое, а задача найти k из Q=kP — дискретный логарифм. В расширенном калькуляторе эти понятия представлены отдельными командами.")
    heading(doc, "Контрольные вопросы для самопроверки", 2)
    add_para(doc, "Вопросы сформулированы для объяснения выполненной реализации; отдельный список вопросов преподавателем не предоставлен.")
    for n, (q, a) in enumerate([
        ("Почему модуль должен быть простым?", "В Fp каждый ненулевой элемент имеет обратный. При составном модуле кольцо вычетов может содержать делители нуля и обычные формулы сложения не обеспечивают заявленную группу."),
        ("Для чего нужна точка O?", "Она является нейтральным элементом, результатом суммы противоположных точек и результатом нулевого скалярного умножения."),
        ("Чем отличается порядок точки от числа точек?", "Порядок точки — размер порождённой ею циклической подгруппы. Число точек — размер всей группы. Первый делит второй, но может быть меньше."),
        ("Почему нельзя делить координаты обычным целочисленным делением?", "Деление в поле — умножение на модульный обратный элемент. Например, 1/2 mod 17 равно 9, потому что 2·9≡1."),
        ("Почему double-and-add эффективнее последовательного сложения?", "Он обрабатывает двоичные разряды скаляра и требует порядка log|k| групповых операций вместо порядка |k|."),
        ("Почему в ECDH совпадают результаты участников?", "A=aG и B=bG, следовательно aB=(ab)G=(ba)G=bA. Общая точка затем требует отдельного преобразования в сеансовый ключ в полноценном протоколе."),
    ], 1):
        add_qa(doc, n, q, a)

    add_page_break(doc)
    heading(doc, "Выводы")
    add_para(doc, "Выполнена практическая работа по построению консольного калькулятора на эллиптических кривых. Реализованы модульная арифметика, проверка параметров, закон сложения и двоичное скалярное умножение, перечисление точек и исследование группы. На конкретных примерах продемонстрированы дискретный логарифм и совпадение общих точек ECDH.")
    add_para(doc, "Результаты вычислений совпали с ручными примерами и независимым эталоном на Python. Автоматические проверки подтвердили корректность особых случаев, свойств группы и арифметики у границы u64. Подготовлены инструкции запуска, воспроизводимые журналы и отчёт с фрагментами фактических исходников.")
    heading(doc, "Список использованных источников", 2)
    add_para(doc, "1. Certicom Research. Standards for Efficient Cryptography 1 (SEC 1): Elliptic Curve Cryptography. Version 2.0, 2009. § 2.2.1. URL: https://www.secg.org/sec1-v2.pdf.")
    add_para(doc, "2. McGrew D., Igoe K., Salter M. Fundamental Elliptic Curve Cryptography Algorithms. RFC 6090, 2011. §§ 2–4. URL: https://www.rfc-editor.org/rfc/rfc6090.")
    add_para(doc, "3. Izykowski W., Worley S. Deterministic variants of the Miller–Rabin primality test. Набор семи оснований Jim Sinclair для n<2⁶⁴. URL: https://miller-rabin.appspot.com/.")
    # Служебные заголовки и подписи не должны оставаться одиночной строкой страницы.
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
    for table_ in doc.tables:
        for row in table_.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
    path = OUT / "Ковалев Д.П. ВКБ43 Практика 2.docx"
    save(doc, path)
    return path


def main():
    for folder in [OUT, EVIDENCE, SHOTS, SNIPS]:
        folder.mkdir(parents=True, exist_ok=True)
    print("Проверка workspace: just ci", flush=True)
    quality = command(["just", "ci"])
    quality_text = quality.stdout + quality.stderr
    (EVIDENCE / "just-ci.log").write_text(quality_text, encoding="utf-8")
    command(["cargo", "build", "-p", "practice_02"])
    tests = command(["cargo", "test", "-p", "practice_02"])
    test_text = tests.stdout + tests.stderr
    (EVIDENCE / "practice-tests.log").write_text(test_text, encoding="utf-8")
    passed = sum(map(int, re.findall(r"test result: ok\. (\d+) passed", test_text)))
    reference = reference_verification()
    metadata = {"checked_at": datetime.now().astimezone().isoformat(timespec="seconds"),
                "rustc": command(["rustc", "--version"]).stdout.strip(),
                "practice_tests": passed, "reference_cases": reference["passed"],
                "just_ci_exit_code": quality.returncode}
    (EVIDENCE / "checks.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Практика 2: {passed} теста; Python: {reference['passed']} сверки", flush=True)
    shots = {
        "01_arithmetic": capture("01_arithmetic", [(args, 0) for args in [
            ["double", "5,1"], ["add", "5,1", "6,3"], ["sub", "5,1", "6,3"],
            ["mul", "7", "5,1"], ["mul", "--", "-3", "5,1"], ["mul", "0", "5,1"],
            ["linear", "2", "5,1", "3", "6,3"], ["add", "5,1", "5,16"],
        ]], max_lines=30, compact_curve=True),
        "02_points": capture("02_points", [(["points"], 0)]),
        "03_subgroup": capture("03_subgroup", [(["order", "5,1"], 0), (["subgroup", "5,1"], 0)], max_lines=30),
        "04_table": capture("04_table", [(["--p", "5", "--a", "1", "--b", "0", "table"], 0)]),
        "05_ecdh": capture("05_ecdh", [(["dlog", "5,1", "0,6"], 0), (["ecdh", "5,1", "5", "7"], 0)]),
        "06_errors": capture("06_errors", [(["--p", "15", "info"], 1), (["--a", "0", "--b", "0", "info"], 1),
            (["check", "1,1"], 0), (["add", "1,1", "O"], 1), (["--p", "10007", "points"], 1)], max_lines=32, summarize_errors=True),
        "07_large": capture("07_large", [(["--p", "18446744073709551557", "--a", "1", "--b", "1", "mul", "123456789", "0,1"], 0)]),
    }
    summary = (f"$ just ci\nfmt-check: OK\nClippy --workspace --all-targets --all-features: OK\n"
               f"cargo test --workspace --all-features: OK\nКод завершения: 0\n\n"
               f"$ cargo test -p practice_02\n{passed} passed; 0 failed\n\n"
               f"Независимая сверка Python: {reference['passed']} passed\n"
               f"Проверено: {metadata['checked_at']}\n")
    (EVIDENCE / "08_quality.log").write_text(summary, encoding="utf-8")
    shots["08_quality"] = render_term_chunked(summary, SHOTS, "08_quality", max_lines=25)
    source = "crates/practice_02/src/domain/"
    snippets = {
        "field": snippet("01_field", source + "field.rs", "    pub fn add(", "    /// Малая теорема"),
        "addition": snippet("02_addition", source + "curve.rs", "    fn add_valid(", "    pub fn subtract("),
        "multiply": snippet("03_multiply", source + "curve.rs", "    pub fn multiply(", "    /// Все точки"),
    }
    report = build_document(metadata, shots, snippets)
    print(report, flush=True)


if __name__ == "__main__":
    main()
