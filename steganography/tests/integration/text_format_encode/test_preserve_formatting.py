"""DOCX covers retain their formatting and package contents during encoding."""

import asyncio
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from steganography.application.commands.text_format_encode.encode import EncodeSecretCommand
from steganography.domain.common.value_objects.formatting_param import FormattingParam
from steganography.infrastructure.text_format_encode.docx_cover_text_reader import DocxCoverTextReaderImpl
from tests.integration.text_format_encode.test_encode_decode_roundtrip import _encode_handler, _decode_handler
from steganography.application.commands.text_format_decode.decode import DetectSecretCommand


@pytest.mark.parametrize("param,zero,one", [
    (FormattingParam.SPACING, "0", "2"),
    (FormattingParam.SCALE, "100", "99"),
    (FormattingParam.COLOR, "000000", "010000"),
    (FormattingParam.SIZE, "28", "29"),
    (FormattingParam.HIGHLIGHT, "yellow", "green"),
])
def test_docx_cover_preserves_formatting(tmp_path: Path, param, zero, one) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "encoded.docx"
    document = Document()
    document.styles["Normal"].font.name = "Cambria"
    document.styles["Normal"].font.size = Pt(13)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(17)
    first = paragraph.add_run("Верба зацвела весной. " * 10)
    first.font.name = "Arial"
    first.font.size = Pt(18)
    first.bold = True
    first.add_tab()
    first.add_break()
    first.add_text("Текст после переноса.")
    second = paragraph.add_run("Ветер дует. " * 10)
    second.font.name = "Times New Roman"
    second.font.size = Pt(11)
    second.italic = True
    document.add_paragraph("Наследуемый стиль.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Таблица"
    document.sections[0].header.paragraphs[0].text = "Колонтитул"
    document.save(source)
    cover = DocxCoverTextReaderImpl().read(source)
    result = asyncio.run(_encode_handler()(EncodeSecretCommand(
        secret_text="Ветер дует.", cover=cover, encoding_name="Windows-1251",
        param=param, zero_value=zero, one_value=one, output_path=output,
    )))
    assert result.success, result.error
    encoded = Document(output)
    assert encoded.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert encoded.paragraphs[0].paragraph_format.space_after == Pt(17)
    assert encoded.paragraphs[0].text == paragraph.text
    visible = [run for run in encoded.paragraphs[0].runs if run.text.strip()]
    assert visible[0].font.name == "Arial"
    assert visible[0].bold is True
    assert visible[-1].font.name == "Times New Roman"
    assert visible[-1].italic is True
    if param is not FormattingParam.SIZE:
        assert visible[0].font.size == Pt(18)
        assert visible[-1].font.size == Pt(11)
        assert encoded.paragraphs[1].runs[0].font.size is None
    assert encoded.tables[0].cell(0, 0).text == "Таблица"
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name)
    decoded = asyncio.run(_decode_handler()(DetectSecretCommand(docx_path=output)))
    assert decoded.success, decoded.error
    assert decoded.message == "Ветер дует."
    assert decoded.method.param is param
