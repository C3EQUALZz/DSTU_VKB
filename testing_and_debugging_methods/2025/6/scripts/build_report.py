"""Отчёт №6: реальные сессии pdb и снимки графического отладчика PyCharm."""

from pathlib import Path
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT.parent / "scripts"))
from report_builder import (LabMeta, add_heading, add_image, add_listing, add_page_break,
                            add_para, add_table_simple, add_title_page, make_doc, save)

SCREENS = ROOT.parent / "docs/screenshots/lab_06"
OUTPUT = ROOT.parent / "docs/reports/lab_06/Ковалев Д.П. ВКБ43 6 лаба.docx"


def main():
    doc = make_doc()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    doc.core_properties.title = "Лабораторная работа №6. Отладка в pdb и PyCharm"
    doc.core_properties.author = "Ковалев Данил Петрович"
    add_title_page(doc, LabMeta(6, "Отладка программы с помощью pdb и графического отладчика PyCharm"))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def page(title):
        add_heading(doc, title)
        doc.paragraphs[-1].paragraph_format.page_break_before = True

    def para(text):
        add_para(doc, text)

    def image(name, caption, width=15.5):
        add_image(doc, SCREENS / (name + ".png"), caption, width_cm=width)
        doc.paragraphs[-2].paragraph_format.keep_with_next = True
        doc.paragraphs[-1].paragraph_format.keep_together = True

    def table(rows, caption):
        add_table_simple(doc, rows, caption=caption)
        t = doc.tables[-1]
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    page("1. Цель работы и исходная программа")
    para("Цель работы: освоить отладку Python-программы в консольном pdb и графическом отладчике PyCharm: установить точки останова, выполнить код по шагам, исследовать переменные и стек вызовов, исправить найденную ошибку и проверить результат.")
    para("По заданию допускается любой небольшой скрипт. Подготовлен учебный расчёт среднего арифметического с намеренно внесённой ошибкой границы цикла. Исходная версия сохранена в average_legacy.py, исправленная — в average.py; demo.py запускает обе версии на одинаковых данных.")
    table([
        ["Параметр", "Значение"],
        ["Входной список", "[10, 20, 30, 40]"],
        ["Контрольный расчёт", "(10 + 20 + 30 + 40) / 4 = 25,00"],
        ["Результат исходной версии", "15,00 вместо 25,00"],
        ["Среда", "macOS; Python 3.14.3; PyCharm 2026.2.2"],
    ], "Таблица 1 — исходные данные и контрольный результат")
    image("01_reproduce", "Рисунок 1 — воспроизводимое расхождение исходной версии с эталоном")
    para("Расчёт не вызывает исключения, поэтому отсутствие traceback не означает правильность программы. Ключ --check выявляет неверный результат и возвращает код завершения 1. Проверяемые гипотезы: неверная граница цикла, ошибка накопления суммы, неверный делитель.")

    page("2. Отладка в pdb: остановка и выполнение по шагам")
    para("Запуск: .venv/bin/python -m pdb demo.py --version legacy. Команда b average_legacy.py:9 задаёт точку перед циклом, c продолжает выполнение до неё. p выводит значения, n выполняет следующую строку без входа во вложенные вызовы, display отслеживает изменение выражения [1].")
    image("02_pdb_steps", "Рисунок 2 — фрагмент реальной сессии pdb: p, display и два шага n")
    para("В начале total равен 0.0. После входа в цикл index равен 0; следующий шаг выполняет сложение и total становится 10.0. Выражение list(range(len(values) - 1)) уже показывает подозрительную границу: [0, 1, 2] при четырёх входных значениях.")
    para("Подсвечиваемая строка отладчика ещё не выполнена. Поэтому при остановке на total += values[index] значение total относится к состоянию до текущего сложения. Полные протоколы сохраняются локально в report_assets/logs/.")

    page("3. Отладка в pdb: локализация причины")
    para("Точка b average_legacy.py:11 останавливает программу перед возвратом среднего. После c проверены исходный список, последний индекс, накопленная сумма, длина списка и независимая сумма через sum(values).")
    image("03_pdb_cause", "Рисунок 3 — перед return получены index=2 и total=60.0")
    table([
        ["Проверка", "Получено", "Вывод"],
        ["p values", "[10, 20, 30, 40]", "Входные данные не потеряны."],
        ["p (index, total, len(values))", "(2, 60.0, 4)", "Обработаны только первые три элемента."],
        ["p values[-1] / p sum(values)", "40 / 100", "Пропущено слагаемое 40."],
        ["p total / len(values)", "15.0", "Делитель 4 верен; неверен числитель."],
    ], "Таблица 2 — наблюдения, различающие гипотезы")
    para("Гипотеза неверной границы цикла подтверждена. Накопление первых трёх значений выполнено верно: 10 + 20 + 30 = 60. Деление также корректно для полученной суммы: 60 / 4 = 15. Команда where дополнительно показала цепочку demo.py → main() → average(); стек включён в полный протокол сессии.")

    page("4. PyCharm: подготовка графической отладки")
    para("В PyCharm создана конфигурация Python «Lab 06 - legacy»: скрипт demo.py, параметры --version legacy --check, рабочий каталог лабораторной №6. Для выполнения выбран Python 3.14.3. Конфигурация запускается через Run → Debug.")
    para("В average_legacy.py установлены точки останова на строках 10 и 11. В панели Threads & Variables отображаются локальные значения; в Watches добавлены len(values), sum(values), list(range(len(values) - 1)) и values[-1]. Стек кадров расположен слева [2].")
    image("04_pycharm_step", "Рисунок 4 — реальная остановка PyCharm на первой итерации цикла", width=16.5)
    para("На первой итерации index=0 и total=0.0. Исходный список содержит четыре значения. Step Over выполняет сложение без входа во внутренние операции; Resume Program продолжает работу до следующей точки останова.")

    page("5. PyCharm: исследование состояния перед return")
    para("Выполнение продолжено до точки на строке 11. В редакторе выделена следующая исполняемая строка return total / len(values). Одновременно видны исходный код, стек, Watches и локальные переменные.")
    image("05_pycharm_bug", "Рисунок 5 — PyCharm подтверждает пропуск последнего элемента", width=16.5)
    table([
        ["Объект наблюдения", "Значение на снимке"],
        ["index / total", "2 / 60.0"],
        ["len(values) / sum(values)", "4 / 100"],
        ["list(range(len(values) - 1))", "[0, 1, 2]"],
        ["values[-1]", "40"],
    ], "Таблица 3 — состояние в графическом отладчике")
    para("Наблюдения совпали с pdb. Для списка длины 4 цикл range(3) проходит индексы 0, 1, 2; индекс 3 не достигается. Это ошибка на единицу в верхней границе диапазона. Изменять входные данные или делитель для её устранения не требуется.")

    page("6. Исправление и регрессионная проверка")
    para("Исправлена строка цикла: range(len(values) - 1) заменено на range(len(values)). Диапазон теперь включает индекс последнего элемента. Формула сложения и деление на длину списка оставлены прежними.")
    image("06_fix", "Рисунок 6 — исправление верхней границы диапазона", width=14.5)
    para("До исправления создан тест test_regression_includes_last_element: среднее для [10, 20, 30, 40] должно равняться 25.0. На исходной строке этот тест обнаружил расхождение 15.0 и 25.0.")
    image("11_regression_before", "Рисунок 7 — регрессионный тест падает на версии до исправления")
    para("После исправления тот же тест проходит. Исходная версия сохранена отдельно; протокол «до» воспроизводится на временном снимке.")

    page("7. Повторная проверка исправления в pdb")
    para("Исправленная версия запущена той же командой с параметром --version fixed. Точка b average.py:11 установлена перед return. Проверяются те же переменные и выражения.")
    image("07_pdb_fixed", "Рисунок 8 — pdb: обработан индекс 3, сумма 100.0, среднее 25.0")
    para("Теперь (index, total, len(values)) = (3, 100.0, 4), а диапазон индексов — [0, 1, 2, 3]. После c программа выводит Average: 25.00; expected: 25.00. Pdb сообщает о нормальном sys.exit(0), предлагает повторный запуск; q завершает сессию.")

    page("8. Повторная проверка исправления в PyCharm")
    para("Конфигурация «Lab 06 - fixed» запускает demo.py с --version fixed --check. Точка на строке 11 файла average.py показывает состояние после обработки всех элементов списка.")
    image("08_pycharm_fixed", "Рисунок 9 — PyCharm: исправленный цикл, index=3 и total=100.0", width=16.5)
    para("В Watches сохранено выражение старой границы для сравнения; выполняемый цикл на строке 9 теперь использует range(len(values)). После Resume Program консоль подтверждает правильный итог и успешное сравнение.")
    image("09_pycharm_result", "Рисунок 10 — фрагмент консоли PyCharm: Average=25.00 и CHECK: PASS", width=16.5)

    page("9. Итоговое тестирование и выводы")
    para("Полный набор pytest содержит 12 тестовых случаев. Проверяются регрессия на последнем элементе, воспроизводимость legacy-дефекта, список из одного значения, два значения, ненулевой последний элемент, отрицательные и дробные значения, нули, одинаковые значения, пустой список и сохранность входных данных.")
    image("10_tests", "Рисунок 11 — полный запуск pytest: 12 passed")
    para("Оба отладчика установили одну причину: последний элемент пропускался из-за range(len(values) - 1). Исправление восстановило сумму 100.0 и среднее 25.0. Повторные остановки перед return подтвердили правильность промежуточного состояния, а тесты — корректность результата на контрольных и граничных примерах.")
    para("Предотвращению ошибки служат тесты одиночного списка и ненулевого последнего элемента. В работе сохранена минимальная правка одной строки.")

    page("10. Сравнение инструментов и воспроизведение")
    table([
        ["Действие", "pdb", "PyCharm"],
        ["Точка останова", "b файл:строка", "Щелчок по полю слева от строки"],
        ["Следующая строка", "n", "Step Over"],
        ["Продолжить", "c", "Resume Program"],
        ["Просмотр значений", "p / display", "Threads & Variables / Watches"],
        ["Стек вызовов", "where", "Панель кадров слева"],
        ["Завершить", "q", "Stop / нормальное завершение"],
    ], "Таблица 4 — соответствие использованных действий")
    add_heading(doc, "Команды из каталога лабораторной №6", level=2)
    add_listing(doc,
        "uv sync --all-groups\n"
        ".venv/bin/python demo.py --version legacy --check\n"
        ".venv/bin/python -m pdb demo.py --version legacy\n"
        "# Внутри pdb: b average_legacy.py:11; c; p total; q\n"
        ".venv/bin/python -m pdb demo.py --version fixed\n"
        ".venv/bin/python demo.py --version fixed --check\n"
        ".venv/bin/python -m pytest\n"
        ".venv/bin/python scripts/capture_evidence.py --stage before\n"
        ".venv/bin/python scripts/capture_evidence.py --stage after\n"
        ".venv/bin/python scripts/build_report.py")
    para("Скриншоты PyCharm сделаны в настоящем окне IDE; их нельзя получить одним запуском capture_evidence.py. Этот скрипт воспроизводит консольные сессии pdb и тесты. Иллюстрации хранятся в docs/screenshots/lab_06, отчёты — в docs/reports/lab_06; локальные протоколы report_assets исключены из Git.")
    add_heading(doc, "Источники", level=2)
    para("1. Python Software Foundation. pdb — The Python Debugger. https://docs.python.org/3/library/pdb.html (дата обращения: 20.10.2026).")
    para("2. JetBrains. Part 1. Debugging Python Code. https://www.jetbrains.com/help/pycharm/part-1-debugging-python-code.html (дата обращения: 20.10.2026).")
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.runs and all(r.bold for r in paragraph.runs) and paragraph.text.strip():
            paragraph.paragraph_format.keep_with_next = True
    save(doc, OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
